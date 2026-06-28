from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MAX_PREVIEW_ROWS = 500
MAX_XLSX_UNCOMPRESSED_SIZE = 50 * 1024 * 1024
MAX_XLSX_ENTRIES = 2_000
MONTHS = {
    "enero": 1,
    "febrero": 2,
    "marzo": 3,
    "abril": 4,
    "mayo": 5,
    "junio": 6,
    "julio": 7,
    "agosto": 8,
    "septiembre": 9,
    "setiembre": 9,
    "octubre": 10,
    "noviembre": 11,
    "diciembre": 12,
}


class ImportFormatError(ValueError):
    pass


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normalized(value: Any) -> str:
    text = unicodedata.normalize("NFKD", _text(value))
    return "".join(char for char in text if not unicodedata.combining(char)).lower()


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    text = _text(value).replace("$", "").replace(" ", "")
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    else:
        text = text.replace(",", "")
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _issue(
    row: int,
    column: str,
    code: str,
    severity: str,
    message: str,
    value: Any,
    suggestion: Any = None,
) -> dict[str, Any]:
    return {
        "row": row,
        "source_year_block": None,
        "column": column,
        "code": code,
        "severity": severity,
        "message": message,
        "value": value,
        "suggestion": suggestion,
    }


def _metadata_from_title(title: str) -> dict[str, str | None]:
    normalized = re.sub(r"\s+", " ", title).strip()
    match = re.search(
        r"precio\s+por\s+(?P<presentation>\w+)\s+de\s+"
        r"(?P<product>.+?)(?:\s+calidad\s+(?P<quality>.+))?$",
        normalized,
        re.IGNORECASE,
    )
    if not match:
        return {
            "title": normalized,
            "product_name": None,
            "quality": None,
            "presentation": None,
        }
    return {
        "title": normalized,
        "product_name": _text(match.group("product")).lower(),
        "quality": _text(match.group("quality")) or None,
        "presentation": _text(match.group("presentation")).lower(),
    }


def _status_for_row(
    row_number: int,
    source_year_block: int | None,
    issues: list[dict[str, Any]],
) -> str:
    severities = {
        issue["severity"]
        for issue in issues
        if issue["row"] == row_number
        and issue.get("source_year_block") == source_year_block
    }
    if "error" in severities:
        return "error"
    if "warning" in severities:
        return "warning"
    return "valid"


def _build_result(
    filename: str,
    sheet: str,
    source_format: str,
    metadata: dict[str, Any],
    rows: list[dict[str, Any]],
    issues: list[dict[str, Any]],
) -> dict[str, Any]:
    for row in rows:
        row["status"] = _status_for_row(
            row["source_row"], row["source_year_block"], issues
        )

    counts = Counter(row["status"] for row in rows)
    return {
        "filename": filename,
        "sheet": sheet,
        "format": source_format,
        "metadata": metadata,
        "summary": {
            "total_rows": len(rows),
            "valid_rows": counts["valid"],
            "warning_rows": counts["warning"],
            "error_rows": counts["error"],
            "issues": len(issues),
        },
        "rows": rows[:MAX_PREVIEW_ROWS],
        "_all_rows": rows,
        "issues": issues,
        "truncated": len(rows) > MAX_PREVIEW_ROWS,
    }


def _parse_historical_workbook(
    content: bytes,
    filename: str,
    producer_name: str | None,
    municipality: str | None,
    currency: str,
    package_weight_kg: Decimal | None,
) -> dict[str, Any]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise ImportFormatError(
                    "El archivo Excel contiene demasiados elementos internos."
                )
            if sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_SIZE:
                raise ImportFormatError(
                    "El contenido descomprimido del Excel excede 50 MB."
                )
            if any(entry.filename.lower().endswith("vbaproject.bin") for entry in entries):
                raise ImportFormatError("No se permiten archivos Excel con macros.")
    except zipfile.BadZipFile as exc:
        raise ImportFormatError("El archivo XLSX está dañado o no es válido.") from exc

    workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    worksheet = workbook.active

    title = next(
        (
            _text(cell.value)
            for row in worksheet.iter_rows(min_row=1, max_row=min(5, worksheet.max_row))
            for cell in row
            if isinstance(cell.value, str) and cell.value.strip()
        ),
        "",
    )
    detected = _metadata_from_title(title)
    metadata = {
        **detected,
        "producer_name": producer_name,
        "municipality": municipality,
        "currency": currency.upper(),
        "package_weight_kg": float(package_weight_kg) if package_weight_kg else None,
    }

    year_blocks: list[tuple[int, int]] = []
    for row in worksheet.iter_rows(min_row=1, max_row=min(8, worksheet.max_row)):
        for cell in row:
            if isinstance(cell.value, (int, float)) and 2000 <= int(cell.value) <= 2100:
                year_blocks.append((int(cell.value), cell.column))
    year_blocks = list(dict.fromkeys(year_blocks))
    if not year_blocks:
        raise ImportFormatError(
            "No se encontraron bloques anuales en las primeras ocho filas."
        )

    rows: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    seen_dates: dict[tuple[int, date], list[int]] = {}

    for block_year, start_column in year_blocks:
        first_issue_for_block = len(issues)
        for source_row in range(6, worksheet.max_row + 1):
            values = [
                worksheet.cell(source_row, start_column + offset).value
                for offset in range(4)
            ]
            if all(value is None for value in values):
                continue

            raw_date, raw_week, raw_month, raw_price = values
            record_date = _date(raw_date)
            price = _decimal(raw_price)
            month_name = _normalized(raw_month)
            week = int(raw_week) if isinstance(raw_week, (int, float)) else None

            if record_date is None:
                issues.append(_issue(
                    source_row, "record_date", "invalid_date", "error",
                    "La fecha no tiene un formato reconocido.", raw_date,
                ))
            if price is None:
                issues.append(_issue(
                    source_row, "price", "invalid_price", "error",
                    "El precio no es numérico.", raw_price,
                ))
            elif price <= 0:
                issues.append(_issue(
                    source_row, "price", "non_positive_price", "error",
                    "El precio debe ser mayor que cero.", float(price), None,
                ))
            elif -price.as_tuple().exponent > 2:
                issues.append(_issue(
                    source_row, "price", "excessive_precision", "warning",
                    "El precio tiene más de dos decimales.",
                    float(price), float(price.quantize(Decimal("0.01"))),
                ))
            if not detected["product_name"]:
                issues.append(_issue(
                    source_row, "product_name", "missing_product", "error",
                    "No se pudo identificar el producto desde el título.",
                    title, None,
                ))

            if record_date:
                if record_date.year != block_year:
                    suggestion = record_date.replace(year=block_year)
                    issues.append(_issue(
                        source_row, "record_date", "year_mismatch", "error",
                        f"La fecha no pertenece al bloque {block_year}.",
                        record_date.isoformat(), suggestion.isoformat(),
                    ))
                expected_month = MONTHS.get(month_name)
                if expected_month and expected_month != record_date.month:
                    issues.append(_issue(
                        source_row, "month", "month_mismatch", "warning",
                        "El mes escrito no coincide con la fecha.",
                        raw_month, record_date.month,
                    ))
                iso_week = record_date.isocalendar().week
                if week is not None and week != iso_week:
                    issues.append(_issue(
                        source_row, "week", "week_mismatch", "warning",
                        "La semana no coincide con la semana ISO de la fecha.",
                        week, iso_week,
                    ))
                duplicate_key = (block_year, record_date)
                previous_rows = seen_dates.setdefault(duplicate_key, [])
                if previous_rows:
                    issues.append(_issue(
                        source_row, "record_date", "duplicate_date", "warning",
                        f"La fecha también aparece en las filas {previous_rows}.",
                        record_date.isoformat(), None,
                    ))
                previous_rows.append(source_row)

            rows.append({
                "source_row": source_row,
                "source_year_block": block_year,
                "record_date": record_date.isoformat() if record_date else None,
                "year": record_date.year if record_date else block_year,
                "iso_week": record_date.isocalendar().week if record_date else week,
                "month": record_date.month if record_date else MONTHS.get(month_name),
                "product_name": detected["product_name"],
                "quality": detected["quality"],
                "presentation": detected["presentation"],
                "price": float(price) if price is not None else None,
                "currency": currency.upper(),
            })
        for issue in issues[first_issue_for_block:]:
            issue["source_year_block"] = block_year

    return _build_result(
        filename, worksheet.title, "annual_blocks", metadata, rows, issues
    )


HEADER_ALIASES = {
    "fecha": "record_date",
    "date": "record_date",
    "precio": "price",
    "precio por caja": "price",
    "price": "price",
    "producto": "product_name",
    "product": "product_name",
    "calidad": "quality",
    "presentacion": "presentation",
    "presentación": "presentation",
    "productor": "producer_name",
    "nombre productor": "producer_name",
    "municipio": "municipality",
    "moneda": "currency",
}


def _parse_csv(
    content: bytes,
    filename: str,
    producer_name: str | None,
    municipality: str | None,
    currency: str,
    package_weight_kg: Decimal | None,
    column_mapping: dict[str, str] | None = None,
) -> dict[str, Any]:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        decoded = content.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(decoded), dialect=dialect)
    if not reader.fieldnames:
        raise ImportFormatError("El CSV no contiene encabezados.")

    mapping = column_mapping or {
        header: HEADER_ALIASES.get(_normalized(header))
        for header in reader.fieldnames
    }
    mapped_fields = {value for value in mapping.values() if value}
    if not {"record_date", "price"} <= mapped_fields:
        raise ImportFormatError(
            "El CSV necesita al menos columnas reconocibles de fecha y precio."
        )

    rows: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    for source_row, raw in enumerate(reader, start=2):
        normalized = {
            target: raw.get(source)
            for source, target in mapping.items()
            if target
        }
        record_date = _date(normalized.get("record_date"))
        price = _decimal(normalized.get("price"))
        if record_date is None:
            issues.append(_issue(
                source_row, "record_date", "invalid_date", "error",
                "La fecha no tiene un formato reconocido.",
                normalized.get("record_date"),
            ))
        if price is None or price <= 0:
            issues.append(_issue(
                source_row, "price", "invalid_price", "error",
                "El precio debe ser un número mayor que cero.",
                normalized.get("price"),
            ))
        if not _text(normalized.get("product_name")):
            issues.append(_issue(
                source_row, "product_name", "missing_product", "error",
                "No se identificó el producto de esta fila.",
                normalized.get("product_name"),
            ))
        rows.append({
            "source_row": source_row,
            "source_year_block": None,
            "record_date": record_date.isoformat() if record_date else None,
            "year": record_date.year if record_date else None,
            "iso_week": record_date.isocalendar().week if record_date else None,
            "month": record_date.month if record_date else None,
            "product_name": normalized.get("product_name"),
            "quality": normalized.get("quality"),
            "presentation": normalized.get("presentation"),
            "price": float(price) if price is not None else None,
            "currency": _text(normalized.get("currency")) or currency.upper(),
        })

    metadata = {
        "title": None,
        "product_name": None,
        "quality": None,
        "presentation": None,
        "producer_name": producer_name,
        "municipality": municipality,
        "currency": currency.upper(),
        "package_weight_kg": float(package_weight_kg) if package_weight_kg else None,
        "column_mapping": mapping,
    }
    return _build_result(filename, "CSV", "flat_table", metadata, rows, issues)


def parse_import_file(
    content: bytes,
    filename: str,
    producer_name: str | None = None,
    municipality: str | None = None,
    currency: str = "MXN",
    package_weight_kg: Decimal | None = None,
    column_mapping: dict[str, str] | None = None,
) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        return _parse_historical_workbook(
            content, filename, producer_name, municipality, currency,
            package_weight_kg,
        )
    if suffix == ".csv":
        return _parse_csv(
            content, filename, producer_name, municipality, currency,
            package_weight_kg, column_mapping,
        )
    raise ImportFormatError("Solo se aceptan archivos .xlsx y .csv.")
