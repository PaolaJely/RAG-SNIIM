import unittest
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from import_parser import ImportFormatError, parse_import_file


ROOT = Path(__file__).resolve().parents[2]


class ImportParserTests(unittest.TestCase):
    def test_annual_blocks_workbook(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Precios"
        worksheet.merge_cells("A1:D1")
        worksheet["A1"] = "PRECIO POR CAJA DE BANANO CALIDAD RC"
        worksheet.merge_cells("A3:D3")
        worksheet["A3"] = 2025
        worksheet.append([])
        worksheet["A4"] = "FECHA"
        worksheet["D4"] = "precio por"
        worksheet["B5"] = "SEMANA"
        worksheet["C5"] = "MES"
        worksheet["D5"] = "caja"
        worksheet.append([date(2025, 1, 3), 1, "enero", 180])
        worksheet.append([date(2026, 1, 10), 2, "enero", 0])
        output = BytesIO()
        workbook.save(output)

        result = parse_import_file(output.getvalue(), "precios.xlsx")

        self.assertEqual(result["format"], "annual_blocks")
        self.assertEqual(result["summary"]["total_rows"], 2)
        self.assertEqual(result["summary"]["valid_rows"], 1)
        self.assertEqual(result["summary"]["error_rows"], 1)
        self.assertEqual(result["metadata"]["product_name"], "banano")
        self.assertEqual(result["metadata"]["quality"], "RC")
        self.assertEqual(result["metadata"]["presentation"], "caja")
        self.assertEqual(
            {issue["code"] for issue in result["issues"]},
            {"non_positive_price", "year_mismatch"},
        )

    @unittest.skipUnless(
        (ROOT / "docs" / "analisis de precios por año.xlsx").exists(),
        "El libro de ejemplo local no está disponible.",
    )
    def test_real_annual_workbook(self):
        path = ROOT / "docs" / "analisis de precios por año.xlsx"
        result = parse_import_file(
            path.read_bytes(),
            path.name,
            producer_name="Productor prueba",
            municipality="Municipio prueba",
            package_weight_kg=18,
        )

        self.assertEqual(result["format"], "annual_blocks")
        self.assertEqual(result["summary"]["total_rows"], 275)
        self.assertEqual(result["summary"]["valid_rows"], 257)
        self.assertEqual(result["summary"]["warning_rows"], 13)
        self.assertEqual(result["summary"]["error_rows"], 5)
        self.assertEqual(result["metadata"]["product_name"], "banano")
        self.assertEqual(result["metadata"]["quality"], "RC")
        self.assertEqual(result["metadata"]["presentation"], "caja")

        codes = {issue["code"] for issue in result["issues"]}
        self.assertTrue({
            "duplicate_date",
            "excessive_precision",
            "non_positive_price",
            "week_mismatch",
            "year_mismatch",
        } <= codes)

    def test_flat_csv(self):
        content = (
            "Fecha;Producto;Calidad;Presentación;Precio\n"
            "01/02/2026;Banano;RC;Caja;125.50\n"
            "02/02/2026;Banano;RC;Caja;0\n"
        ).encode()
        result = parse_import_file(content, "productor.csv")

        self.assertEqual(result["format"], "flat_table")
        self.assertEqual(result["summary"]["total_rows"], 2)
        self.assertEqual(result["summary"]["valid_rows"], 1)
        self.assertEqual(result["summary"]["error_rows"], 1)
        self.assertEqual(result["rows"][0]["record_date"], "2026-02-01")
        self.assertEqual(result["rows"][0]["price"], 125.5)

    def test_flat_csv_with_agent_mapping(self):
        content = (
            "Día de corte;Valor comercial;Cultivo\n"
            "01/02/2026;125.50;Banano\n"
        ).encode()
        result = parse_import_file(
            content,
            "productor.csv",
            column_mapping={
                "Día de corte": "record_date",
                "Valor comercial": "price",
                "Cultivo": "product_name",
            },
        )

        self.assertEqual(result["summary"]["valid_rows"], 1)
        self.assertEqual(result["rows"][0]["product_name"], "Banano")

    def test_rejects_unknown_extension(self):
        with self.assertRaises(ImportFormatError):
            parse_import_file(b"test", "datos.xls")


if __name__ == "__main__":
    unittest.main()
